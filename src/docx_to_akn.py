"""Convert IPBES Plenary session reports from Word (.docx) to Akoma Ntoso
XML (AKN4UN), in the form of the reports in data/xml.

    PYTHONPATH=src python src/docx_to_akn.py                    # data/docx/*.docx
    PYTHONPATH=src python src/docx_to_akn.py data/docx/2017-5.docx
    PYTHONPATH=src python src/docx_to_akn.py --out data/xml_compare

Files are named <year>-<session>.docx; the XML keeps the name. Old Word
files (.doc) are skipped: open them in Word and save them as .docx first.

The model follows the AKN4UN guidelines for meeting reports
(https://unsceb-hlcm.github.io/part2/, section 7) and naming convention
(https://unsceb-hlcm.github.io/part1/, section 6), as the Indigo files in
data/xml do:

    <akomaNtoso><debateReport name="debateReport">
      <meta>        FRBR identification: /akn/un-ipbes/debatereport/report/plenary/<first day>/<session>
      <preface>     title, session, place and dates, document symbol
      <mainBody>    division (I., II.) > subdivision (A., B.) > paragraph (1., 2.) > subparagraph ((a), (i))
      <attachments> Annex > decision IPBES-5/1 > Annex I to decision IPBES-5/1 > Appendix
                    each attachment: <heading>, <subheading>, <doc name="annexure|decision|appendix">

The values that belong to IPBES (FRBR country and IRIs, the organization in
<references>, the acronym in decision symbols "IPBES-5/1", the document
symbol, Word style names) are in config/akn4un_ipbes.toml. To convert
another body's reports, copy it, change the values and pass --config.

The Word reports use the UN templates: headings are styled CH1-CH4 with the
number typed before a tab ("\\tI.\\tOpening of the session"), paragraphs are
numbered by Word ("Normal_number"), and the annexes start with the
ZZ_Anx_header style. Word's automatic numbers are not in the text; they are
recomputed from numbering.xml (see Numbering).

Compare the result with data/xml using xml_compare.py.
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
import tomllib
from dataclasses import dataclass, field
from pathlib import Path

import docx
from lxml import etree

AKN_NS = "http://docs.oasis-open.org/legaldocml/ns/akn/3.0"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ROOT = Path(__file__).resolve().parent.parent
DEFAULT_CONFIG = ROOT / "config" / "akn4un_ipbes.toml"

# The organization's settings (config/akn4un_ipbes.toml): metadata, the
# acronym in decision symbols, Word styles. Set by load_config().
CONFIG: dict = {}


def w(tag: str) -> str:
    return f"{{{W_NS}}}{tag}"


def akn(tag: str) -> str:
    return f"{{{AKN_NS}}}{tag}"


# ---------------------------------------------------------------------------
# Text
# ---------------------------------------------------------------------------

PAGE_BREAK = "\f"
# [text] double_quotes = "straight": “The Future We Want” -> "The Future We Want"
_STRAIGHT_QUOTES = str.maketrans({"\u201c": '"', "\u201d": '"', "\u201e": '"'})

# Non-breaking hyphens (U+2011) become "-", so "IPBES-5/1" matches searches.
_CHAR_MAP = str.maketrans({"‑": "-", "­": "", "​": ""})


@dataclass
class Span:
    """A run of text with its formatting, or a footnote (note is set)."""
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    vert: str = ""                      # "sup", "sub" or ""
    note: list | None = None            # footnote: list of Span lists (its paragraphs)


def _on(rpr, tag: str) -> bool:
    """A run property such as <w:b/> or <w:b w:val="0"/>."""
    el = rpr.find(w(tag)) if rpr is not None else None
    return el is not None and el.get(w("val"), "true") not in ("0", "false", "none")


def _inside(el, tags: set[str], stop) -> bool:
    parent = el.getparent()
    while parent is not None and parent is not stop:
        if parent.tag in tags:
            return True
        parent = parent.getparent()
    return False


_SKIP_RUNS = {w("del"), w("moveFrom"), w("txbxContent")}


def spans_of(p_el, footnotes: dict | None = None) -> list[Span]:
    """The text of a Word paragraph, as formatted spans, with its footnotes."""
    spans: list[Span] = []
    for r in p_el.iter(w("r")):
        if _inside(r, _SKIP_RUNS, p_el):
            continue
        rpr = r.find(w("rPr"))
        fmt = dict(bold=_on(rpr, "b"), italic=_on(rpr, "i"), underline=_on(rpr, "u"))
        va = rpr.find(w("vertAlign")) if rpr is not None else None
        if va is not None:
            fmt["vert"] = {"superscript": "sup", "subscript": "sub"}.get(va.get(w("val")), "")
        for child in r:
            tag = child.tag
            if tag == w("t"):
                spans.append(Span(child.text or "", **fmt))
            elif tag == w("tab"):
                spans.append(Span("\t", **fmt))
            elif tag in (w("br"), w("cr")):
                # A page break inside a paragraph starts a new one ("…gaps.
                # [page] Annex III to decision IPBES-4/1"): see read_blocks.
                spans.append(Span(PAGE_BREAK if child.get(w("type")) == "page" else "\n", **fmt))
            elif tag == w("noBreakHyphen"):
                spans.append(Span("-", **fmt))
            elif tag == w("sym"):
                spans.append(Span(chr(int(child.get(w("char"), "20"), 16)) if child.get(w("char"), "").isalnum() else "", **fmt))
            elif tag == w("footnoteReference") and footnotes is not None:
                note = footnotes.get(child.get(w("id")))
                if note:
                    spans.append(Span("", note=note))
    quotes = _STRAIGHT_QUOTES if CONFIG.get("text", {}).get("double_quotes") == "straight" else {}
    for s in spans:
        s.text = s.text.translate(_CHAR_MAP).translate(quotes)
    return spans


def plain(spans: list[Span]) -> str:
    return "".join(s.text for s in spans)


def strip_spans(spans: list[Span], prefix_len: int) -> list[Span]:
    """Drop the first `prefix_len` characters (a typed number), then leading
    and trailing white space."""
    out, skip = [], prefix_len
    for s in spans:
        if skip and s.note is None:
            cut = min(skip, len(s.text))
            skip -= cut
            s = Span(s.text[cut:], s.bold, s.italic, s.underline, s.vert)
        out.append(s)
    # leading / trailing white space
    while out and out[0].note is None and not out[0].text.strip():
        out.pop(0)
    if out and out[0].note is None:
        out[0] = Span(out[0].text.lstrip(), out[0].bold, out[0].italic, out[0].underline, out[0].vert)
    while out and out[-1].note is None and not out[-1].text.strip():
        out.pop()
    if out and out[-1].note is None:
        out[-1] = Span(out[-1].text.rstrip(), out[-1].bold, out[-1].italic, out[-1].underline, out[-1].vert)
    return out


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


# ---------------------------------------------------------------------------
# Numbering: Word's automatic paragraph and list numbers
# ---------------------------------------------------------------------------

def _roman(n: int) -> str:
    out = ""
    for value, letters in ((1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"),
                           (90, "xc"), (50, "l"), (40, "xl"), (10, "x"), (9, "ix"),
                           (5, "v"), (4, "iv"), (1, "i")):
        while n >= value:
            out, n = out + letters, n - value
    return out


def _letter(n: int) -> str:
    """1 -> a, 26 -> z, 27 -> aa (Word repeats the letter)."""
    return chr(ord("a") + (n - 1) % 26) * ((n - 1) // 26 + 1)


def format_number(n: int, fmt: str) -> str:
    return {
        "lowerLetter": _letter(n), "upperLetter": _letter(n).upper(),
        "lowerRoman": _roman(n), "upperRoman": _roman(n).upper(),
        "bullet": "", "none": "",
    }.get(fmt, str(n))


@dataclass
class Level:
    start: int = 1
    fmt: str = "decimal"
    text: str = "%1."


class Numbering:
    """Recomputes the numbers Word shows for numbered paragraphs.

    Each list (w:num) points to a list definition (w:abstractNum) with a
    format per level. Lists sharing a definition share its counters; a list
    with a w:startOverride restarts them where it is first used. A number
    at one level restarts the levels below it.
    """

    def __init__(self, document):
        self.abstract: dict[str, dict[int, Level]] = {}
        self.nums: dict[str, tuple[str, dict[int, int], dict[int, Level]]] = {}
        self.style_num: dict[str, tuple[str | None, int | None]] = {}
        self.style_parent: dict[str, str] = {}
        self.counters: dict = {}
        self.used: set = set()
        self._read_styles(document)
        try:
            root = document.part.numbering_part.element
        except (NotImplementedError, KeyError, AttributeError):
            return
        self.style_links: dict[str, str] = {}      # abstractNum -> list style it uses
        for a in root.findall(w("abstractNum")):
            levels = {}
            for lvl in a.findall(w("lvl")):
                levels[int(lvl.get(w("ilvl")))] = self._level(lvl)
            self.abstract[a.get(w("abstractNumId"))] = levels
            link = a.find(w("numStyleLink"))
            if link is not None:
                self.style_links[a.get(w("abstractNumId"))] = link.get(w("val"))
        for n in root.findall(w("num")):
            starts, overrides = {}, {}
            for o in n.findall(w("lvlOverride")):
                ilvl = int(o.get(w("ilvl")))
                so = o.find(w("startOverride"))
                if so is not None:
                    starts[ilvl] = int(so.get(w("val")))
                lvl = o.find(w("lvl"))
                if lvl is not None:
                    overrides[ilvl] = self._level(lvl)
            abs_id = n.find(w("abstractNumId")).get(w("val"))
            self.nums[n.get(w("numId"))] = (abs_id, starts, overrides)

    @staticmethod
    def _level(lvl) -> Level:
        def val(tag, default):
            el = lvl.find(w(tag))
            return el.get(w("val")) if el is not None and el.get(w("val")) is not None else default
        return Level(int(val("start", "1")), val("numFmt", "decimal"), val("lvlText", ""))

    def _read_styles(self, document):
        for style in document.styles.element.findall(w("style")):
            sid = style.get(w("styleId"))
            based = style.find(w("basedOn"))
            if based is not None:
                self.style_parent[sid] = based.get(w("val"))
            numpr = style.find(f"{w('pPr')}/{w('numPr')}")
            if numpr is not None:
                num_id = numpr.find(w("numId"))
                ilvl = numpr.find(w("ilvl"))
                self.style_num[sid] = (num_id.get(w("val")) if num_id is not None else None,
                                       int(ilvl.get(w("val"))) if ilvl is not None else None)

    def _resolve_link(self, abs_id: str) -> str:
        """A list definition that only points to a list style
        (w:numStyleLink) takes its levels from the list style's numbering."""
        seen = set()
        while abs_id in self.style_links and abs_id not in seen:
            seen.add(abs_id)
            num_id, _ = self.style_num.get(self.style_links[abs_id], (None, None))
            if num_id not in self.nums:
                break
            abs_id = self.nums[num_id][0]
        return abs_id

    def _style_numpr(self, style_id: str | None):
        seen = set()
        while style_id and style_id not in seen:
            seen.add(style_id)
            if style_id in self.style_num:
                return self.style_num[style_id]
            style_id = self.style_parent.get(style_id)
        return None, None

    def number(self, p_el) -> tuple[str, str, int] | None:
        """(number text, format, level) for a numbered paragraph, else None.
        Advances the list's counters, so call it once per paragraph, in
        document order."""
        ppr = p_el.find(w("pPr"))
        num_id = ilvl = None
        style_id = None
        if ppr is not None:
            ps = ppr.find(w("pStyle"))
            style_id = ps.get(w("val")) if ps is not None else None
            numpr = ppr.find(w("numPr"))
            if numpr is not None:
                n = numpr.find(w("numId"))
                i = numpr.find(w("ilvl"))
                num_id = n.get(w("val")) if n is not None else None
                ilvl = int(i.get(w("val"))) if i is not None else None
        if num_id is None:
            s_num, s_ilvl = self._style_numpr(style_id)
            num_id = s_num
            ilvl = ilvl if ilvl is not None else s_ilvl
        if not num_id or num_id == "0" or num_id not in self.nums:
            return None
        ilvl = ilvl or 0
        abs_id, starts, overrides = self.nums[num_id]
        abs_id = self._resolve_link(abs_id)
        levels = dict(self.abstract.get(abs_id, {}))
        levels.update(overrides)
        if ilvl not in levels:
            return None
        counts = self.counters.setdefault(abs_id, {})
        first_use = (num_id, ilvl) not in self.used
        self.used.add((num_id, ilvl))
        if ilvl in starts and first_use:
            counts[ilvl] = starts[ilvl]         # the list restarts the numbering
        else:
            counts[ilvl] = counts[ilvl] + 1 if ilvl in counts else levels[ilvl].start
        for deeper in [k for k in counts if k > ilvl]:
            del counts[deeper]
        level = levels[ilvl]
        if level.fmt in ("bullet", "none"):
            return "", level.fmt, ilvl

        def fill(m):
            k = int(m.group(1)) - 1
            lv = levels.get(k, Level())
            value = counts.get(k, starts.get(k, lv.start))
            return format_number(value, lv.fmt)
        return re.sub(r"%(\d)", fill, level.text).strip(), level.fmt, ilvl


# ---------------------------------------------------------------------------
# Reading the Word file into a flat list of blocks
# ---------------------------------------------------------------------------

# Word styles, from the settings ([styles]); see load_config().
HEADING_STYLES: dict[str, int] = {}
TITLE_STYLES: set[str] = set()
ANNEX_HEADER_STYLES: set[str] = set()
ANNEX_TITLE_STYLES: set[str] = set()
SKIP_STYLES: set[str] = set()

# A number typed before the text: "I.", "A.", "1.", "(a)", "(i)", "a)".
_TYPED_NUMBER = re.compile(
    r"^[\t ]*(\(?[0-9]{1,3}[.)]|\([a-zA-Z]{1,4}\)|[a-z]{1,2}\)|[IVXLC]{1,6}\.|[A-Z]\.|[0-9]{1,2}\.[0-9]{1,2}\.?)"
    r"(?:\t+| {2,})")
# Draft alternatives in negotiated texts: "[2bis.", "[2bis alt", "3ter."
# (IPBES-1 rules of procedure), followed by a single space.
_BIS_NUMBER = re.compile(r"^[\t ]*(\[?\d{1,3}\s?(?:bis|ter|quater)(?:\s+alt)?\.?)(?=\s)")
_HEADING_NUMBER = re.compile(r"^[\t ]*([IVXLC]{1,6}\.?|[A-Z]\.|[0-9]{1,2}(?:\.[0-9]{1,2})*\.?)\t+")
_ROMAN_ONLY = re.compile(r"^[IVXL]{1,6}\.?$")


@dataclass
class Block:
    kind: str                       # heading, para, table, title, annex_header, annex_title
    spans: list = field(default_factory=list)
    style: str = ""
    level: float = 0                # heading level
    num: str = ""                   # number (Word's or typed)
    fmt: str = ""                   # decimal, lowerLetter, lowerRoman, upperRoman, bullet ...
    ilvl: int = 0
    auto: bool = False              # number from Word's numbering
    rows: list = field(default_factory=list)   # table: rows of cells (dict)

    @property
    def text(self) -> str:
        return norm(plain(self.spans))

    @property
    def bold_line(self) -> bool:
        """A short, unnumbered paragraph in bold: a heading typed as text."""
        return (self.kind == "para" and not self.num and len(self.text) < 300
                and any(s.text.strip() for s in self.spans)
                and all(s.bold for s in self.spans if s.note is None and s.text.strip()))


def _footnotes(document, numbering) -> dict:
    """footnote id -> its paragraphs (lists of spans)."""
    notes = {}
    for rel in document.part.rels.values():
        if rel.reltype.endswith("/footnotes"):
            root = etree.fromstring(rel.target_part.blob)
            for fn in root.findall(w("footnote")):
                if fn.get(w("type")) in ("separator", "continuationSeparator", "continuationNotice"):
                    continue
                paras = [strip_spans(spans_of(p), 0) for p in fn.iter(w("p"))]
                paras = [p for p in paras if plain(p).strip()]
                if paras:
                    notes[fn.get(w("id"))] = paras
    return notes


def _cell_blocks(tc, footnotes) -> list[list[Span]]:
    return [s for s in (strip_spans(spans_of(p, footnotes), 0) for p in tc.iter(w("p"))) if plain(s).strip()]


def _table(tbl, footnotes) -> list[list[dict]]:
    """Rows of cells {paras, colspan, rowspan}; vertically merged cells
    become one cell with a rowspan."""
    rows, origin = [], {}           # origin: grid column -> cell that started a vertical merge
    for tr in tbl.findall(w("tr")):
        col, row = 0, []
        trpr = tr.find(w("trPr"))
        gb = trpr.find(w("gridBefore")) if trpr is not None else None
        if gb is not None:
            col += int(gb.get(w("val"), "0"))
        for tc in tr.findall(w("tc")):
            tcpr = tc.find(w("tcPr"))
            span_el = tcpr.find(w("gridSpan")) if tcpr is not None else None
            span = int(span_el.get(w("val"), "1")) if span_el is not None else 1
            vm = tcpr.find(w("vMerge")) if tcpr is not None else None
            if vm is not None and vm.get(w("val"), "continue") == "continue" and col in origin:
                origin[col]["rowspan"] += 1
                col += span
                continue
            cell = {"paras": _cell_blocks(tc, footnotes), "colspan": span, "rowspan": 1}
            if vm is not None:
                origin[col] = cell
            else:
                origin.pop(col, None)
            row.append(cell)
            col += span
        if row:
            rows.append(row)
    return rows


def _body_elements(parent):
    """Paragraphs and tables of the body in order, looking inside content
    controls (w:sdt)."""
    for el in parent:
        if el.tag in (w("p"), w("tbl")):
            yield el
        elif el.tag == w("sdt"):
            content = el.find(w("sdtContent"))
            if content is not None:
                yield from _body_elements(content)


def read_blocks(path: Path) -> list[Block]:
    document = docx.Document(str(path))
    numbering = Numbering(document)
    footnotes = _footnotes(document, numbering)
    styles = {s.style_id: s.name for s in document.styles}
    blocks: list[Block] = []
    for el in _body_elements(document.element.body):
        if el.tag == w("tbl"):
            blocks.append(Block("table", rows=_table(el, footnotes)))
            continue
        ppr = el.find(w("pPr"))
        ps = ppr.find(w("pStyle")) if ppr is not None else None
        style = styles.get(ps.get(w("val")) if ps is not None else "Normal", "Normal")
        number = numbering.number(el)
        if style in SKIP_STYLES:
            continue
        pieces = _split_pages(spans_of(el, footnotes))
        for k, spans in enumerate(pieces):
            b = _paragraph_block(spans, style, number if k == 0 else None)
            if b is not None:
                blocks.append(b)
    return blocks


def _split_pages(spans: list[Span]) -> list[list[Span]]:
    """A paragraph's spans, split where it has a page break."""
    pieces: list[list[Span]] = [[]]
    for s in spans:
        if s.note is None and PAGE_BREAK in s.text:
            parts = s.text.split(PAGE_BREAK)
            for n, part in enumerate(parts):
                if n:
                    pieces.append([])
                if part:
                    pieces[-1].append(Span(part, s.bold, s.italic, s.underline, s.vert))
        else:
            pieces[-1].append(s)
    return [p for p in pieces if plain(p).strip() or any(s.note for s in p)]


def _paragraph_block(spans: list[Span], style: str, number) -> "Block | None":
    """A block for one paragraph (or the part of one after a page break)."""
    raw = plain(spans)
    if not raw.strip() and not any(s.note for s in spans):
        return None
    b = Block("para", style=style)
    if style in HEADING_STYLES:
        b.kind, b.level = "heading", HEADING_STYLES[style]
    elif style in TITLE_STYLES:
        b.kind = "title"
    elif style in ANNEX_HEADER_STYLES:
        b.kind = "annex_header"
    elif style in ANNEX_TITLE_STYLES:
        b.kind = "annex_title"
    prefix = 0
    if number is not None:
        b.num, b.fmt, b.ilvl = number
        b.auto = True
    else:
        m = (_HEADING_NUMBER if b.kind == "heading" else _TYPED_NUMBER).match(raw) \
            or (_BIS_NUMBER.match(raw) if b.kind == "para" else None)
        if m:
            b.num, prefix = m.group(1), m.end()
            b.fmt = _typed_format(b.num)
    b.spans = strip_spans(spans, prefix)
    if b.kind == "annex_header" and not re.match(rf"\[?\s*(Annex|Appendi|Decision|{_BODY})", b.text, re.I):
        b.kind = "para"                 # a table note in the annex header style
    return b if b.text or any(s.note for s in b.spans) else None


def _typed_format(num: str) -> str:
    if re.search(r"\d\s?(?:bis|ter|quater)", num):
        return "decimal"
    core = num.strip("().")
    if core.isdigit() or re.fullmatch(r"[0-9.]+", core):
        return "decimal"
    if re.fullmatch(r"[IVXLC]+", core) and core not in ("C", "L"):
        return "upperRoman"
    if re.fullmatch(r"[ivxlc]+", core) and len(core) > 1:
        return "lowerRoman"
    if re.fullmatch(r"[A-Z]", core):
        return "upperLetter"
    if re.fullmatch(r"[ivxlc]", core):
        return "lowerRoman?"            # (i) or the letter i: decided in context
    return "lowerLetter"


# ---------------------------------------------------------------------------
# Building the document tree
# ---------------------------------------------------------------------------

# Decision symbols ("IPBES-5/1"; the acronym is [document] body in the
# settings), compiled by load_config():
_BODY = "IPBES"
_DECISION: re.Pattern            # "IPBES-5/1: ...", "Decision IPBES-7/1", "Decision IPBES/1/2"
_MISTYPED_DECISION: re.Pattern   # typed differently in a heading: "IPBES 4-4" (IPBES-4)
_ANNEX_TO_DECISION: re.Pattern   # "Annex I to decision IPBES-5/1"
_OLD_DECISION: re.Pattern        # "Decision IPBES/1/2 Next steps ..." (IPBES-1)
_SYMBOL: re.Pattern              # document symbol on the cover: IPBES/5/15
_ANNEX = re.compile(r"^Annex(?:es)?\b", re.IGNORECASE)
_TABLE_TITLE = re.compile(r"^(Table\s+[0-9]+[A-Za-z]?)\.?(?:\s|$)")
_APPENDIX = re.compile(r"^Appendi(?:x|ces)\b", re.IGNORECASE)


def load_config(path: Path = DEFAULT_CONFIG) -> dict:
    """Read the organization's settings and set the styles and patterns
    that depend on them."""
    global _BODY, _DECISION, _MISTYPED_DECISION, _ANNEX_TO_DECISION, _OLD_DECISION, _SYMBOL
    with open(path, "rb") as f:
        CONFIG.clear()
        CONFIG.update(tomllib.load(f))
    styles = CONFIG["styles"]
    HEADING_STYLES.clear()
    HEADING_STYLES.update(styles["headings"])
    for target, key in ((TITLE_STYLES, "titles"), (ANNEX_HEADER_STYLES, "annex_headers"),
                        (ANNEX_TITLE_STYLES, "annex_titles"), (SKIP_STYLES, "skip")):
        target.clear()
        target.update(styles.get(key, []))
    _BODY = re.escape(CONFIG["document"]["body"])
    _DECISION = re.compile(rf"^(?:Decision\s+)?{_BODY}[-‑/]\d+/\d+\b", re.IGNORECASE)
    _MISTYPED_DECISION = re.compile(rf"^((?:Decision\s+)?){_BODY}[\s-]?(\d+)[-‑](\d+)\b", re.IGNORECASE)
    _ANNEX_TO_DECISION = re.compile(rf"\bto\s+decision\s+({_BODY}[-‑]\d+/\d+)", re.IGNORECASE)
    _OLD_DECISION = re.compile(rf"^Decision\s+({_BODY}/\d+/\d+)\s*:?\s*")
    _SYMBOL = re.compile(CONFIG["document"]["symbol"])
    return CONFIG

# Rank of a list format: lower ranks contain higher ones.
_FMT_RANK = {"decimal": 1, "upperLetter": 1, "lowerLetter": 2, "lowerRoman": 3, "upperRoman": 1, "bullet": 4}


@dataclass
class Node:
    tag: str                        # division, subdivision, paragraph, subparagraph, attachment, root
    num: str = ""
    heading: list = field(default_factory=list)       # spans
    subheading: list = field(default_factory=list)
    level: float = 0                # heading level (hierarchy) or list rank (paragraphs)
    intro: list = field(default_factory=list)         # blocks before children
    children: list = field(default_factory=list)      # Nodes
    wrap: list = field(default_factory=list)          # blocks after children
    kind: str = ""                  # attachment: top, decision, annex, appendix
    decision: str = ""              # attachment: decision symbol (IPBES-5/1)
    report_annex: bool = False      # attachment: a decision that is also an annex of the report (IPBES-1)
    body: "Node | None" = None      # attachment: its document root

    def add_block(self, block):
        (self.wrap if self.children else self.intro).append(block)


class Builder:
    """Places blocks into the tree: attachments > headings > paragraphs >
    subparagraphs."""

    def __init__(self):
        self.root = Node("root")
        self.attachments: list[Node] = []          # open attachments, outermost first
        self.preface: list[Block] = []
        self.started = False
        self._reset_doc(self.root)
        self.pending_subheading = False

    def _reset_doc(self, doc_root: Node):
        self.doc_root = doc_root
        self.hier: list[Node] = []                  # open headings
        self.para: Node | None = None               # open paragraph
        self.items: list[Node] = []                 # open subparagraphs
        self.last_para_num: int | None = None       # number of the last paragraph

    def container(self) -> Node:
        return self.hier[-1] if self.hier else self.doc_root

    # attachments -----------------------------------------------------------
    def open_attachment(self, kind: str, heading: list):
        text = norm(plain(heading)).lstrip("[ ")
        decision = ""
        if kind == "decision":
            decision = _DECISION.match(text).group(0).split()[-1]
        stack = self.attachments
        if kind == "top":
            stack.clear()
        elif kind == "decision":
            while stack and stack[-1].kind != "top":
                stack.pop()
        elif kind == "annex":
            target = _ANNEX_TO_DECISION.search(text)
            # "Annex II" after "Annex I - Decision IPBES/1/1" is the next
            # annex of the report, not an annex of that decision.
            in_decision = any(a.kind == "decision" and not a.report_annex for a in stack)
            if target or in_decision:
                while stack and stack[-1].kind != "decision":
                    stack.pop()
            else:
                stack.clear()
        elif kind == "appendix":
            while stack and stack[-1].kind not in ("annex", "decision"):
                stack.pop()
        att = Node("attachment", heading=heading, kind=kind, decision=decision)
        att.body = Node("root")
        (stack[-1].body if stack else self.root).children.append(att) if stack else self.root.children.append(att)
        stack.append(att)
        self._reset_doc(att.body)
        self.pending_subheading = True

    # headings --------------------------------------------------------------
    def open_heading(self, level: float, num: str, heading: list):
        self.close_table_section()
        while self.hier and self.hier[-1].level >= level:
            self.hier.pop()
        # Top-level headings are divisions, the levels below subdivisions,
        # as in most of the reference files.
        depth = len(self.hier)
        node = Node("division" if depth == 0 else "subdivision", num=num,
                    heading=heading, level=level)
        self.container().children.append(node)
        self.hier.append(node)
        self.para, self.items = None, []

    def open_table_section(self, num: str, heading: list):
        """A table title ("Table 1 / Status of cash contributions ...") is a
        subdivision holding the table, as in the reference files: num
        "Table 1", eId ...__subdvs_Table1__table_1."""
        self.close_table_section()
        top = self.hier[-1] if self.hier else None
        node = Node("subdivision", num=num, heading=heading, kind="table",
                    level=(top.level if top is not None else 0) + 0.25)
        self.container().children.append(node)
        self.hier.append(node)
        self.para, self.items = None, []

    def close_table_section(self):
        if self.hier and self.hier[-1].kind == "table":
            self.hier.pop()

    # paragraphs ------------------------------------------------------------
    def add_paragraph(self, b: Block):
        self.close_table_section()
        m = re.match(r"\(?(\d+)", b.num)
        self.last_para_num = int(m.group(1)) if m else None
        node = Node("paragraph", num=b.num, intro=[b])
        self.container().children.append(node)
        self.para, self.items = node, []

    def add_item(self, b: Block, rank: float):
        if self.para is None:
            self.close_table_section()
        # A list without a numbered paragraph: its items are paragraphs
        # ("(a)", "(b)" after a heading), siblings at the same level.
        if self.para is not None and 0 < rank <= self.para.level:
            self.para, self.items = None, []
        if self.para is None:
            node = Node("paragraph", num=b.num, intro=[b], level=rank)
            self.container().children.append(node)
            self.para, self.items = node, []
            return
        while self.items and self.items[-1].level >= rank:
            self.items.pop()
        parent = self.items[-1] if self.items else self.para
        node = Node("subparagraph", num=b.num, intro=[b], level=rank)
        parent.children.append(node)
        self.items.append(node)

    def add_text(self, b: Block):
        """An unnumbered block: it continues the open paragraph."""
        if self.para is not None:
            target = self.para
            if self.items:
                target = self.para          # text after a list closes the list
                self.items = []
            target.add_block(b)
        else:
            self.container().add_block(b)


def _list_rank(b: Block, previous_letter: str | None) -> float:
    fmt = b.fmt
    if fmt == "lowerRoman?":
        core = b.num.strip("().")
        fmt = "lowerLetter" if previous_letter and ord(core) == ord(previous_letter) + 1 else "lowerRoman"
        b.fmt = fmt
    rank = _FMT_RANK.get(fmt, 2)
    # The format decides: 1. contains (a) contains (i). Word's list level
    # only orders items of the same format (1. then 1.1), since items from
    # different lists have unrelated levels (agenda "2." at level 3, its
    # "(a)" at level 1).
    return rank + 0.1 * b.ilvl if b.auto else rank


def build(blocks: list[Block]) -> Builder:
    bld = Builder()
    i = 0
    letter_before: str | None = None
    last_letter: str | None = None              # letter of the last "A." heading
    while i < len(blocks):
        b = blocks[i]
        nxt = blocks[i + 1] if i + 1 < len(blocks) else None
        text = b.text
        title = text.lstrip("[ ")               # "[Appendix", "[[Annex IV": bracketed drafts

        # Front matter: everything before the first heading.
        if not bld.started:
            if b.kind == "heading" and not b.num and text.startswith("Report of"):
                b.kind, b.style = "title", "BB_Title"     # the title in a heading style
            if b.kind == "heading" or b.kind == "annex_header":
                bld.started = True
            else:
                bld.preface.append(b)
                i += 1
                continue

        if b.kind == "title":                   # a title style after the cover: a heading
            b.kind, b.level = "heading", 1
        heading_like = b.kind in ("heading", "annex_title", "annex_header") or b.bold_line
        # Attachments: annex header, decisions, annexes, appendices.
        m = _MISTYPED_DECISION.match(title) if heading_like and not b.num else None
        if m:                                   # "IPBES 4-4: ..." -> "IPBES-4/4: ..."
            fixed = f"{m.group(1)}{CONFIG['document']['body']}-{m.group(2)}/{m.group(3)}"
            b.spans = [Span(fixed)] + strip_spans(b.spans, len(plain(b.spans)) - len(plain(b.spans).lstrip()) + m.end())
            text = title = b.text
        if b.kind == "annex_header" and _DECISION.match(title):
            b.kind = "heading"                  # a decision in the annex header style
        if b.kind == "annex_header":
            in_decision = any(a.kind == "decision" for a in bld.attachments)
            bld.open_attachment("annex" if in_decision and _ANNEX.match(title) else "top", b.spans)
            i += 1
            continue
        # "Annex I" then "Decision IPBES/1/1 Rules of procedure ..." (IPBES-1):
        # the decision is the annex's subheading (see _decision_in_subheading).
        if heading_like and not b.num and bld.pending_subheading and _OLD_DECISION.match(title):
            att = bld.attachments[-1] if bld.attachments else None
            if att is not None and att.kind == "annex" and not att.subheading:
                att.subheading = b.spans
                _decision_in_subheading(att)
                i += 1
                continue
        if heading_like and not b.num and _starts_attachment(title):
            kind = ("decision" if _DECISION.match(title) else
                    "appendix" if _APPENDIX.match(title) else "annex")
            bld.open_attachment(kind, b.spans)
            i += 1
            continue
        if heading_like and not b.num and bld.pending_subheading:
            att = bld.attachments[-1] if bld.attachments else None
            if att is not None and not att.subheading:
                att.subheading = b.spans
                _decision_in_subheading(att)
                i += 1
                continue
        bld.pending_subheading = False

        # Headings. An unnumbered heading straight before another heading
        # contains it ("Executive summary" over "A. Rationale").
        # A table title: "Table 1" (num) and its title (heading), in a
        # heading or table title style, in bold, or just before a table.
        tm = _TABLE_TITLE.match(text) if b.kind in ("para", "heading") and not b.num else None
        if tm and len(text) < 250 and not text.endswith(".") and (
                b.kind == "heading" or b.style == "Title_table" or b.bold_line
                or any(x.kind == "table" for x in blocks[i + 1:i + 3])):
            raw = plain(b.spans)
            cut = len(raw) - len(raw.lstrip()) + len(re.match(r"\s*Table\s+\S+?\.?(?=\s|$)", raw).group(0))
            bld.open_table_section(norm(tm.group(1)), strip_spans(b.spans, cut))
            i += 1
            continue

        # A section number alone on its line in a heading style ("VIII"),
        # then the title: one heading, as in decisions.
        if b.kind == "heading" and not b.num and _ROMAN_ONLY.match(text) and nxt is not None \
                and nxt.kind in ("heading", "para") and not nxt.num and len(nxt.text) < 200:
            bld.open_heading(1, text, nxt.spans)
            i += 2
            continue
        if b.kind in ("heading", "annex_title"):
            # The number decides the level when it has one (some reports
            # style "B." as CH3 next to "A." in CH2); the style otherwise.
            core = b.num.strip("().")
            if b.fmt == "upperRoman" and last_letter and len(core) == 1 \
                    and ord(core) == ord(last_letter) + 1:
                b.fmt = "upperLetter"           # "I." after "H." is a letter
            level = {"upperRoman": 1, "upperLetter": 2, "decimal": 3}.get(b.fmt) if b.num else None
            level = level or b.level or 1
            if b.fmt == "upperLetter":
                last_letter = core[-1:]
            elif b.fmt == "upperRoman":
                last_letter = None
            if not b.num and nxt is not None and nxt.kind == "heading":
                level -= 0.5
            bld.open_heading(level, b.num, b.spans)
            i += 1
            continue
        # A section number alone on its line ("II"), then the title: a
        # heading in decisions.
        if not b.auto and _ROMAN_ONLY.match(text) and nxt is not None and nxt.kind == "para" \
                and not nxt.num and len(nxt.text) < 200:
            bld.open_heading(1, text, nxt.spans)
            i += 2
            continue
        # Headings typed as bold paragraphs, in older reports: numbered
        # ("1.1 Governance structures", "II. Initial ...") or not ("Rule 25").
        bold_heading = _bold_heading(b)
        if bold_heading is not None:
            level, num, spans = bold_heading
            if level is None:                   # unnumbered: under the open heading
                top = bld.hier[-1] if bld.hier else None
                level = top.level if top is not None and top.kind == "bold" else \
                    (top.level + 0.5 if top is not None else 1)
            bld.open_heading(level, num, spans)
            if num == "" and bld.hier:
                bld.hier[-1].kind = "bold"
            i += 1
            continue

        if b.kind == "table":
            bld.add_text(b)
        elif b.num and b.fmt != "bullet" and _is_paragraph_number(b, bld):
            bld.add_paragraph(b)
        elif b.num or b.fmt == "bullet":
            rank = _list_rank(b, letter_before)
            bld.add_item(b, rank)
        else:
            bld.add_text(b)
        # The last letter of an (a), (b) list: "(i)" after "(h)" is a letter.
        # Text between the items keeps it ("(b) ... [note] (c)").
        if b.fmt == "lowerLetter":
            letter_before = b.num.strip("().")[-1:]
        elif b.num and b.fmt not in ("lowerRoman", "lowerRoman?"):
            letter_before = None
        i += 1
    return bld


_BOLD_TYPED_NUMBER = re.compile(r"^([IVXL]{1,5}|[A-Z]|\d{1,2}(?:\.\d{1,2})*)\.?\s+(?=\S)")


def _bold_heading(b: Block):
    """(level, number, spans) when a bold paragraph is a heading, else
    None. Level from the number: I. 1, A. 2, 1. 3, 1.1 4, 1.1.1 5;
    None for an unnumbered one."""
    if b.kind != "para" or b.fmt == "bullet":
        return None
    texts = [s for s in b.spans if s.note is None and s.text.strip()]
    if not texts or not all(s.bold for s in texts):
        return None
    text = b.text
    if len(text) > 160 or text.endswith((".", ":", ";", ",")):
        return None
    num, spans = b.num, b.spans
    if not num:
        m = _BOLD_TYPED_NUMBER.match(text)
        if m and not text.startswith(("Rule ", "Table ", "Figure ")):
            num = m.group(1) + ("." if not m.group(1)[-1:].isdigit() or "." not in m.group(1) else "")
            spans = strip_spans(spans, len(plain(spans)) - len(plain(spans).lstrip()) + m.end())
        else:
            return None, "", spans
    core = num.strip("().")
    if re.fullmatch(r"[IVXL]+", core) and core not in ("L",):
        return 1, num, spans
    if re.fullmatch(r"[A-Z]", core):
        return 2, num, spans
    if re.fullmatch(r"\d+(?:\.\d+)*", core):
        return 2 + core.count(".") + 1, num, spans
    return None


def _decision_in_subheading(att: Node):
    """IPBES-1: "Annex III" with the subheading "Decision IPBES/1/2 Next
    steps ...". The decision goes into the heading ("Annex III - Decision
    IPBES/1/2"), as in the Indigo files, so it is read as a decision."""
    text = norm(plain(att.subheading))
    m = _OLD_DECISION.match(text)
    if att.kind == "decision" or not m:
        return
    att.kind, att.decision = "decision", m.group(1)
    att.report_annex = True             # an annex of the report, not of a decision
    att.heading = [Span(f"{norm(plain(att.heading))} - Decision {m.group(1)}")]
    att.subheading = strip_spans(att.subheading, len(plain(att.subheading)) - len(plain(att.subheading).lstrip()) + m.end())


def _starts_attachment(text: str) -> bool:
    return bool(_DECISION.match(text) or _ANNEX.match(text) or _APPENDIX.match(text))


def _is_paragraph_number(b: Block, bld: Builder) -> bool:
    """A numbered paragraph ("12."), as opposed to an item of a list: the
    next number of the paragraphs, or the first number after a heading. A
    "1." inside a paragraph starts a list ("The Plenary adopted the
    following agenda: 1. Opening of the session ...")."""
    if b.fmt != "decimal" or b.ilvl > 0:
        return False
    m = re.match(r"[(\[]?(\d+)", b.num)
    if not m:
        return False
    n = int(m.group(1))
    if re.search(r"bis|ter|quater", b.num):
        return True                     # "[2bis." after paragraph 2
    if bld.para is None:
        return True
    # The next item of an open list ("10." after item "9." of an agenda
    # in paragraph 9).
    if bld.items:
        top = re.match(r"\(?(\d+)", bld.items[0].num)
        if top and int(top.group(1)) + 1 == n:
            return False
    if bld.last_para_num is not None and n == bld.last_para_num + 1:
        return True
    return not (n == 1 or bld.items or _introduces_list(bld))


def _introduces_list(bld: Builder) -> bool:
    if bld.para is None or bld.items or bld.para.children:
        return False
    last = bld.para.wrap[-1] if bld.para.wrap else bld.para.intro[-1] if bld.para.intro else None
    return bool(last is not None and getattr(last, "kind", "") == "para" and last.text.endswith(":"))


# ---------------------------------------------------------------------------
# Metadata from the front matter
# ---------------------------------------------------------------------------

MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august",
          "september", "october", "november", "december"]
_M = "(" + "|".join(m.capitalize() for m in MONTHS) + ")"
_DATE_RANGE = re.compile(
    rf"(\d{{1,2}})(?:\s+{_M})?(?:\s+(\d{{4}}))?\s*(?:–|-|—|to)\s*\d{{1,2}}\s+{_M}\s+(\d{{4}})")
_ORDINALS = ["first", "second", "third", "fourth", "fifth", "sixth", "seventh", "eighth",
             "ninth", "tenth", "eleventh", "twelfth", "thirteenth", "fourteenth", "fifteenth"]


@dataclass
class Info:
    session: str
    date: str               # first day of the session, YYYY-MM-DD
    title: str
    place: str
    symbol: str


def document_info(path: Path, bld: Builder, all_blocks: list[Block]) -> Info:
    front = [b for b in bld.preface if b.kind != "table"]
    front_text = " ".join(b.text for b in front)
    tables = " ".join(" ".join(" ".join(plain(p) for p in c["paras"]) for r in b.rows for c in r)
                      for b in bld.preface if b.kind == "table")
    session = ""
    m = re.match(CONFIG["document"]["file_name"], report_name(path))
    if m:
        session = m.group(2)
    else:
        for n, word in enumerate(_ORDINALS, 1):
            if re.search(rf"\b{word} session\b", front_text, re.IGNORECASE):
                session = str(n)
                break
    title = next((b.text for b in front if b.style == "BB_Title"), "")
    if not title:
        title = next((b.text for b in front if b.text.startswith("Report of")), "")
    if not title:
        title = CONFIG["document"]["title"].format(ordinal=_ORDINALS[int(session) - 1])
    date, place = "", ""
    early = front_text + " " + " ".join(b.text for b in all_blocks[:60] if b.kind == "para")
    dm = _DATE_RANGE.search(early)
    if dm:
        day, month1, year1, month2, year2 = dm.groups()
        month = month1 or month2
        year = year1 or year2
        date = f"{year}-{MONTHS.index(month.lower()) + 1:02d}-{int(day):02d}"
    elif m:
        date = f"{m.group(1)}-01-01"
    for b in front:
        if b.style == "AA_Title" and _DATE_RANGE.search(b.text):
            place = b.text
    symbol_match = _SYMBOL.search(tables + " " + front_text)
    return Info(session, date, title, place, symbol_match.group(0) if symbol_match else "")


# ---------------------------------------------------------------------------
# Writing Akoma Ntoso
# ---------------------------------------------------------------------------

class Writer:
    def __init__(self, info: Info):
        self.info = info
        ident = CONFIG["identification"]
        self.work = (f"/akn/{ident['country']}/{ident['doc_type']}/{ident['subtype']}/"
                     f"{ident['actor']}/{info.date}/{info.session}")
        self.expr = f"{self.work}/{ident['language']}@{info.date}"
        self.used_ids: set[str] = set()
        self.note_number = 0

    # identifiers -----------------------------------------------------------
    def eid(self, prefix: str, kind: str, num: str, index: int) -> str:
        """AKN4UN element identifier: <parent>__<kind>_<num>, the number
        without punctuation ("(a)" -> "a", "I." -> "I"), or a running number."""
        token = re.sub(r"[^0-9A-Za-z]", "", num) if num else ""
        base = f"{kind}_{token or ('nn_' + str(index) if kind in ('dvs', 'subdvs') else index)}"
        value = f"{prefix}__{base}" if prefix else base
        candidate, n = value, 2
        while candidate in self.used_ids:
            candidate, n = f"{value}_{n}", n + 1
        self.used_ids.add(candidate)
        return candidate

    # meta ------------------------------------------------------------------
    def meta(self, component: str, title: str, top: bool) -> etree._Element:
        info = self.info
        today = datetime.date.today().isoformat()
        c, org = CONFIG["identification"], CONFIG["organization"]
        author, ext = c["author_href"], c["manifestation_extension"]
        meta = etree.Element(akn("meta"))
        ident = etree.SubElement(meta, akn("identification"), source=c["source"])
        work = etree.SubElement(ident, akn("FRBRWork"))
        etree.SubElement(work, akn("FRBRthis"), value=f"{self.work}/!{component}")
        etree.SubElement(work, akn("FRBRuri"), value=self.work)
        etree.SubElement(work, akn("FRBRalias"), value=title, name="title")
        if top and info.symbol and c.get("symbol_alias", True):
            etree.SubElement(work, akn("FRBRalias"), value=info.symbol, name="symbol")
        etree.SubElement(work, akn("FRBRdate"), date=info.date, name="Generation")
        etree.SubElement(work, akn("FRBRauthor"), href=author)
        etree.SubElement(work, akn("FRBRcountry"), value=c["country"])
        etree.SubElement(work, akn("FRBRsubtype"), value=c["subtype"])
        etree.SubElement(work, akn("FRBRnumber"), value=info.session)
        expr = etree.SubElement(ident, akn("FRBRExpression"))
        etree.SubElement(expr, akn("FRBRthis"), value=f"{self.expr}/!{component}")
        etree.SubElement(expr, akn("FRBRuri"), value=self.expr)
        etree.SubElement(expr, akn("FRBRdate"), date=info.date, name="Generation")
        etree.SubElement(expr, akn("FRBRauthor"), href=author)
        etree.SubElement(expr, akn("FRBRlanguage"), language=c["language"])
        man = etree.SubElement(ident, akn("FRBRManifestation"))
        etree.SubElement(man, akn("FRBRthis"), value=f"{self.expr}/!{component}{ext}")
        etree.SubElement(man, akn("FRBRuri"), value=f"{self.expr}{ext}")
        etree.SubElement(man, akn("FRBRdate"), date=today, name="Generation")
        etree.SubElement(man, akn("FRBRauthor"), href=author)
        if top:
            refs = etree.SubElement(meta, akn("references"), source=c["source"])
            etree.SubElement(refs, akn("TLCOrganization"), eId=org["eId"], href=org["href"],
                             showAs=org["showAs"])
        return meta

    # inline ----------------------------------------------------------------
    def inline(self, parent: etree._Element, spans: list[Span], eid: str):
        """Formatted text and footnotes, appended to `parent`."""
        last = None

        def add_text(text):
            nonlocal last
            text = text.replace("\t", " ").replace("\n", " ")
            if last is None:
                parent.text = (parent.text or "") + text
            else:
                last.tail = (last.tail or "") + text

        for s in _merge(spans):
            if s.note is not None:
                self.note_number += 1
                n = sum(1 for _ in parent.iter(akn("authorialNote"))) + 1
                note_id = f"{eid}__authorialNote_{n}"
                # In <sup>, as in every reference file.
                sup = etree.SubElement(parent, akn("sup"))
                note = etree.SubElement(sup, akn("authorialNote"), marker=str(self.note_number),
                                        placement="bottom", eId=note_id)
                for k, para in enumerate(s.note, 1):
                    p = etree.SubElement(note, akn("p"), eId=f"{note_id}__p_{k}")
                    self.inline(p, para, f"{note_id}__p_{k}")
                last = sup
                continue
            if not (s.bold or s.italic or s.underline or s.vert):
                add_text(s.text)
                continue
            outer = inner = None
            for flag, tag in ((s.bold, "b"), (s.italic, "i"), (s.underline, "u"),
                              (bool(s.vert), s.vert)):
                if flag:
                    el = etree.Element(akn(tag))
                    if inner is None:
                        outer = el
                    else:
                        inner.append(el)
                    inner = el
            inner.text = s.text.replace("\t", " ").replace("\n", " ")
            parent.append(outer)
            last = outer

    def p(self, parent, spans, prefix: str, index: int):
        eid = f"{prefix}__p_{index}" if prefix else f"p_{index}"
        el = etree.SubElement(parent, akn("p"), eId=eid)
        self.inline(el, spans, eid)
        return el

    def blocks(self, parent, blocks: list[Block], prefix: str):
        """Paragraph text and tables into `parent`."""
        p_count = t_count = 0
        for b in blocks:
            if b.kind == "table":
                t_count += 1
                self.table(parent, b, f"{prefix}__table_{t_count}" if prefix else f"table_{t_count}")
            else:
                p_count += 1
                spans = b.spans
                if b.style in ("Title_table", "Title_figure"):
                    spans = [Span(s.text, True, s.italic, s.underline, s.vert, s.note) for s in spans]
                self.p(parent, spans, prefix, p_count)

    def table(self, parent, b: Block, eid: str):
        table = etree.SubElement(parent, akn("table"), eId=eid)
        k = 0
        for row in b.rows:
            tr = etree.SubElement(table, akn("tr"))
            for cell in row:
                attrs = {}
                if cell["colspan"] > 1:
                    attrs["colspan"] = str(cell["colspan"])
                if cell["rowspan"] > 1:
                    attrs["rowspan"] = str(cell["rowspan"])
                td = etree.SubElement(tr, akn("td"), **attrs)
                for para in cell["paras"] or [[]]:
                    k += 1
                    self.p(td, para, eid, k)
            if len(tr) == 0:
                table.remove(tr)

    # structure -------------------------------------------------------------
    def node(self, parent, node: Node, prefix: str, index: int):
        abbr = {"division": "dvs", "subdivision": "subdvs", "paragraph": "para",
                "subparagraph": "subpara"}[node.tag]
        eid = self.eid(prefix, abbr, node.num, index)
        el = etree.SubElement(parent, akn(node.tag), eId=eid)
        if node.num:
            etree.SubElement(el, akn("num")).text = node.num
        if node.heading:
            self.inline(etree.SubElement(el, akn("heading")), node.heading, eid)
        if node.children:
            if node.intro:
                self.blocks(etree.SubElement(el, akn("intro")), node.intro, f"{eid}__intro")
            for k, child in enumerate(node.children, 1):
                self.node(el, child, eid, k)
            if node.wrap:
                self.blocks(etree.SubElement(el, akn("wrapUp")), node.wrap, f"{eid}__wrapup")
        elif node.intro or node.wrap or node.tag in ("paragraph", "subparagraph"):
            content = etree.SubElement(el, akn("content"))
            self.blocks(content, node.intro + node.wrap, eid)
            if len(content) == 0:
                etree.SubElement(content, akn("p"), eId=f"{eid}__p_1")
        # A heading with nothing under it ("D. Methodological assessment ...",
        # taken up later in the report) is just its number and heading.

    def body(self, parent, root: Node, prefix: str):
        """mainBody of the report or of an attachment; returns the
        attachments found in it."""
        body = etree.SubElement(parent, akn("mainBody"))
        self.blocks(body, root.intro, prefix)
        attachments = []
        for k, child in enumerate((c for c in root.children if c.tag != "attachment"), 1):
            self.node(body, child, prefix, k)
        if root.wrap:
            self.blocks(body, root.wrap, f"{prefix}__wrapup" if prefix else "wrapup")
        attachments = [c for c in root.children if c.tag == "attachment"]
        if len(body) == 0:
            etree.SubElement(body, akn("p"), eId=f"{prefix}__p_1" if prefix else "p_1")
        return attachments

    def attachments(self, parent, atts: list[Node], prefix: str, component: str):
        if not atts:
            return
        container = etree.SubElement(parent, akn("attachments"))
        names = CONFIG["doc_names"]
        per_name: dict[str, int] = {}
        for k, att in enumerate(atts, 1):
            eid = f"{prefix}__att_{k}" if prefix else f"att_{k}"
            self.used_ids.add(eid)
            # <doc name> and component name: annexure_1, annexure_2, appendix_1
            name = names.get(att.kind, names["annex"])
            per_name[name] = per_name.get(name, 0) + 1
            part = f"{name}_{per_name[name]}"
            comp = f"{component}/{part}" if component else part
            el = etree.SubElement(container, akn("attachment"), eId=eid)
            heading = norm(plain(att.heading))
            self.inline(etree.SubElement(el, akn("heading")), att.heading, eid)
            if att.subheading:
                self.inline(etree.SubElement(el, akn("subheading")), att.subheading, eid)
            doc = etree.SubElement(el, akn("doc"), name=name)
            title = " ".join(filter(None, [heading, norm(plain(att.subheading))]))
            doc.append(self.meta(comp, title, top=False))
            inner = self.body(doc, att.body, eid)
            self.attachments(doc, inner, eid, comp)

    def preface(self, parent, bld: Builder):
        info = self.info
        preface = etree.SubElement(parent, akn("preface"))
        if info.symbol:
            p = etree.SubElement(preface, akn("p"), eId="preface__p_1")
            etree.SubElement(p, akn("docNumber")).text = info.symbol
        p = etree.SubElement(preface, akn("p"), eId="preface__p_2")
        etree.SubElement(p, akn("docTitle")).text = info.title
        if info.session:
            p = etree.SubElement(preface, akn("p"), eId="preface__p_3")
            etree.SubElement(p, akn("session"), value=info.session).text = \
                f"{_ORDINALS[int(info.session) - 1].capitalize()} session"
        if info.place:
            p = etree.SubElement(preface, akn("p"), eId="preface__p_4")
            etree.SubElement(p, akn("docDate"), date=info.date).text = info.place

    def document(self, bld: Builder) -> etree._Element:
        root = etree.Element(akn("akomaNtoso"), nsmap={None: AKN_NS})
        report = etree.SubElement(root, akn("debateReport"), name=CONFIG["document"]["name"])
        report.append(self.meta("main", self.info.title, top=True))
        self.preface(report, bld)
        atts = self.body(report, bld.root, "")
        self.attachments(report, atts, "", "")
        return root


def _merge(spans: list[Span]) -> list[Span]:
    """Join neighbouring spans with the same formatting."""
    out: list[Span] = []
    for s in spans:
        if s.note is None and not s.text:
            continue
        if out and s.note is None and out[-1].note is None and \
                (out[-1].bold, out[-1].italic, out[-1].underline, out[-1].vert) == \
                (s.bold, s.italic, s.underline, s.vert):
            out[-1] = Span(out[-1].text + s.text, s.bold, s.italic, s.underline, s.vert)
        else:
            out.append(s)
    # Formatting on white space only is dropped.
    return [Span(s.text, False, False, False, "") if s.note is None and not s.text.strip() else s
            for s in out]


# ---------------------------------------------------------------------------
# Command line
# ---------------------------------------------------------------------------

def report_name(path: Path) -> str:
    """"2016-4.docx" -> "2016-4"; also for a file saved as "2016-4.doc.docx"."""
    return path.stem.removesuffix(".doc")


def convert(path: Path) -> tuple[etree._Element, Info]:
    if not CONFIG:
        load_config()
    blocks = read_blocks(path)
    bld = build(blocks)
    info = document_info(path, bld, blocks)
    return Writer(info).document(bld), info


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("files", nargs="*", type=Path, help="Word files (default: data/docx/*)")
    ap.add_argument("--out", type=Path, default=ROOT / "data" / "xml_compare")
    ap.add_argument("--config", type=Path, default=DEFAULT_CONFIG,
                    help="organization settings (default: config/akn4un_ipbes.toml)")
    args = ap.parse_args(argv)
    load_config(args.config)
    files = args.files or sorted((ROOT / "data" / "docx").glob("*.doc*"))
    args.out.mkdir(parents=True, exist_ok=True)
    for path in files:
        if path.suffix.lower() == ".doc":
            print(f"skip {path.name}: old Word format; save it as .docx in Word")
            continue
        root, info = convert(path)
        out = args.out / f"{report_name(path)}.xml"
        out.write_bytes(etree.tostring(root, xml_declaration=True, encoding="utf-8"))
        counts = {t: sum(1 for _ in root.iter(akn(t))) for t in
                  ("division", "paragraph", "subparagraph", "attachment", "table", "authorialNote")}
        print(f"{path.name} -> {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out}  "
              f"session {info.session}, {info.date}, {info.symbol or 'no symbol'}  "
              + ", ".join(f"{k} {v}" for k, v in counts.items()))


if __name__ == "__main__":
    sys.exit(main())
